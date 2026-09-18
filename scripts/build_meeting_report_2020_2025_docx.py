#!/usr/bin/env python3
"""Build a meeting report DOCX for the 2020-2025 corpus expansion."""

from __future__ import annotations

import csv
import json
import re
from collections import Counter
from datetime import datetime
from pathlib import Path
from urllib.parse import parse_qs, unquote_plus, urlparse

import matplotlib.pyplot as plt
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "artifacts" / "meeting_report_2020_2025"
ASSET_DIR = OUT_DIR / "figures"
OUT_DOCX = OUT_DIR / "2026-06-27--nlp-dataset-discovery--2020-2025-progress-report.docx"


def read_json(path: str) -> dict:
    p = ROOT / path
    if not p.exists():
        return {}
    return json.loads(p.read_text(encoding="utf-8"))


def csv_year_counts(path: str) -> tuple[int, int, dict[str, int]]:
    p = ROOT / path
    rows = 0
    base_ids: set[str] = set()
    years: Counter[str] = Counter()
    if not p.exists():
        return 0, 0, {}
    with p.open(newline="", encoding="utf-8") as handle:
        for row in csv.DictReader(handle):
            rows += 1
            arxiv_id = (row.get("arXiv ID") or "").strip()
            if arxiv_id:
                base_ids.add(arxiv_id.split("v")[0])
            year = (row.get("Publication Date") or "")[:4]
            if year:
                years[year] += 1
    return rows, len(base_ids), dict(sorted(years.items()))


def residual_arxiv_errors() -> list[dict[str, str]]:
    path = ROOT / "artifacts/arxiv_repair_targeted_2020_2022.log"
    if not path.exists():
        return []
    pattern = re.compile(r"Failed arXiv API request for category=(\S+) start_index=(\d+) url=(\S+)")
    errors: list[dict[str, str]] = []
    for line in path.read_text(encoding="utf-8", errors="replace").splitlines():
        match = pattern.search(line)
        if not match:
            continue
        category = match.group(1)
        start_index = match.group(2)
        url = match.group(3)
        query = parse_qs(urlparse(url).query)
        search = unquote_plus(query.get("search_query", [""])[0])
        date_match = re.search(r"submittedDate:\[(\d+) TO (\d+)\]", search)
        if not date_match:
            continue
        errors.append(
            {
                "from": date_match.group(1),
                "to": date_match.group(2),
                "category": category,
                "start": start_index,
            }
        )
    return errors


def save_year_chart(acl_years: dict[str, int], raw_years: dict[str, int], screening_years: dict[str, int]) -> Path:
    years = [str(year) for year in range(2020, 2026)]
    x = range(len(years))
    width = 0.26
    fig, ax = plt.subplots(figsize=(8.2, 4.2))
    ax.bar([i - width for i in x], [acl_years.get(y, 0) for y in years], width, label="ACL catalog", color="#4C78A8")
    ax.bar(x, [raw_years.get(y, 0) for y in years], width, label="arXiv repaired corpus", color="#72B7B2")
    ax.bar([i + width for i in x], [screening_years.get(y, 0) for y in years], width, label="arXiv screening catalog", color="#F58518")
    ax.set_xticks(list(x))
    ax.set_xticklabels(years)
    ax.set_ylabel("Papers / rows")
    ax.set_title("Corpus coverage by year")
    ax.legend(frameon=False, ncol=3, loc="upper left")
    ax.grid(axis="y", color="#DDDDDD", linewidth=0.8)
    ax.spines[["top", "right"]].set_visible(False)
    fig.tight_layout()
    out = ASSET_DIR / "coverage_by_year.png"
    fig.savefig(out, dpi=180)
    plt.close(fig)
    return out


def save_funnel_chart(summary: dict) -> Path:
    labels = ["Repaired corpus", "Title dedup", "Remove ACL overlap", "Screening set"]
    values = [
        int(summary.get("raw_rows") or 0),
        int(summary.get("unique_titles_after_arxiv_dedup") or 0),
        int(summary.get("unique_titles_after_arxiv_dedup") or 0) - int(summary.get("removed_acl_title_overlap") or 0),
        int(summary.get("output_rows") or 0),
    ]
    fig, ax = plt.subplots(figsize=(8.2, 3.5))
    colors = ["#4C78A8", "#72B7B2", "#ECA82C", "#54A24B"]
    ax.barh(labels[::-1], values[::-1], color=colors[::-1])
    for i, value in enumerate(values[::-1]):
        ax.text(value + max(values) * 0.01, i, f"{value:,}", va="center", fontsize=9)
    ax.set_title("arXiv preparation funnel")
    ax.set_xlabel("Rows")
    ax.grid(axis="x", color="#DDDDDD", linewidth=0.8)
    ax.spines[["top", "right", "left"]].set_visible(False)
    fig.tight_layout()
    out = ASSET_DIR / "arxiv_funnel.png"
    fig.savefig(out, dpi=180)
    plt.close(fig)
    return out


def save_status_chart() -> Path:
    stages = [
        ("ACL catalog", 1.0, "Done"),
        ("arXiv repaired corpus", 1.0, "Done"),
        ("Additional arXiv repair", 0.85, "Run done; residual timeouts"),
        ("Final arXiv refresh", 0.25, "Next"),
        ("LLM screening", 0.0, "Not started"),
        ("Fulltext extraction", 0.0, "Not started"),
        ("Metadata enrichment", 0.15, "Smoke ready"),
    ]
    fig, ax = plt.subplots(figsize=(8.2, 4.0))
    y = list(range(len(stages)))
    vals = [s[1] for s in stages]
    colors = ["#54A24B" if v == 1.0 else "#ECA82C" if v > 0 else "#B8B8B8" for v in vals]
    ax.barh(y, vals, color=colors)
    ax.set_yticks(y)
    ax.set_yticklabels([s[0] for s in stages])
    ax.set_xlim(0, 1.0)
    ax.set_xticks([0, 0.5, 1.0])
    ax.set_xticklabels(["Not started", "In progress", "Done"])
    ax.invert_yaxis()
    for i, (_, value, label) in enumerate(stages):
        ax.text(min(value + 0.03, 0.96), i, label, va="center", fontsize=9)
    ax.set_title("Pipeline status")
    ax.grid(axis="x", color="#DDDDDD", linewidth=0.8)
    ax.spines[["top", "right", "left"]].set_visible(False)
    fig.tight_layout()
    out = ASSET_DIR / "pipeline_status.png"
    fig.savefig(out, dpi=180)
    plt.close(fig)
    return out


def save_residual_error_chart(errors: list[dict[str, str]]) -> Path:
    month_counts = Counter(f"{err['from'][:4]}-{err['from'][4:6]}" for err in errors)
    labels = sorted(month_counts)
    values = [month_counts[label] for label in labels]
    fig, ax = plt.subplots(figsize=(8.2, 3.2))
    if labels:
        ax.bar(labels, values, color="#E45756")
        for i, value in enumerate(values):
            ax.text(i, value + 0.05, str(value), ha="center", fontsize=9)
    ax.set_title("Residual arXiv timeout windows after targeted repair")
    ax.set_ylabel("Failed requests")
    ax.set_ylim(0, max(values + [1]) + 1)
    ax.grid(axis="y", color="#DDDDDD", linewidth=0.8)
    ax.spines[["top", "right"]].set_visible(False)
    fig.tight_layout()
    out = ASSET_DIR / "residual_errors.png"
    fig.savefig(out, dpi=180)
    plt.close(fig)
    return out


def set_doc_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    for style_name, size in [("Title", 18), ("Heading 1", 14), ("Heading 2", 12)]:
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)


def add_paragraph(doc: Document, text: str, bold_prefix: str | None = None) -> None:
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(6)
    if bold_prefix and text.startswith(bold_prefix):
        run = para.add_run(bold_prefix)
        run.bold = True
        para.add_run(text[len(bold_prefix):])
    else:
        para.add_run(text)


def add_callout(doc: Document, title: str, lines: list[str]) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    shading = cell._tc.get_or_add_tcPr()
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "EEF3F7")
    shading.append(shd)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    r.bold = True
    for line in lines:
        para = cell.add_paragraph(line)
        para.paragraph_format.space_after = Pt(2)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, header in enumerate(headers):
        run = hdr[i].paragraphs[0].add_run(header)
        run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(2)


def add_figure(doc: Document, path: Path, caption: str, width: float = 6.4) -> None:
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run()
    run.add_picture(str(path), width=Inches(width))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    cap_run = cap.add_run(caption)
    cap_run.italic = True
    cap_run.font.size = Pt(9)


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)

    acl = read_json("data/census/acl_anthology/acl_anthology_2020_2025_all_summary.json")
    screening = read_json("data/processed/arxiv_2020_2025_dedup_no_acl_for_dataset_screening_summary.json")
    raw_rows, raw_base_ids, raw_years = csv_year_counts("data/raw/arxiv_results_2020_2025_repaired.csv")
    targeted_rows, targeted_base_ids, targeted_years = csv_year_counts("data/raw/arxiv_results_2020_2025_repair_targeted.csv")
    errors = residual_arxiv_errors()

    coverage = save_year_chart(acl.get("scope_by_year", {}), raw_years, screening.get("year_counts", {}))
    funnel = save_funnel_chart(screening)
    status = save_status_chart()
    residual = save_residual_error_chart(errors)

    doc = Document()
    set_doc_styles(doc)
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    title = doc.add_paragraph()
    title.style = doc.styles["Title"]
    title.add_run("NLP Dataset Discovery Corpus Expansion Progress Report")
    subtitle = doc.add_paragraph()
    subtitle.add_run("2020-2025 expansion status, open issues, next steps, and cost estimate | 2026-06-27")
    subtitle.paragraph_format.space_after = Pt(12)

    add_callout(
        doc,
        "Bottom line",
        [
            "The 2020-2025 ACL corpus is complete at the paper-catalog level.",
            "The 2020-2025 arXiv corpus has been substantially repaired and prepared for screening.",
            "A small number of arXiv timeout windows remain before treating the arXiv corpus as final.",
        ],
    )

    doc.add_heading("1. Current State", level=1)
    add_table(
        doc,
        ["Component", "Current status", "Current output / count", "Remaining issue"],
        [
            ["ACL corpus", "Complete", f"{acl.get('scope_rows', 0):,} papers", "No major open issue"],
            ["arXiv corpus", "Near final", f"{raw_rows:,} repaired records before final targeted update", "Small residual timeout windows"],
            ["arXiv screening set", "Prepared", f"{screening.get('output_rows', 0):,} papers after deduplication and ACL-overlap removal", "Needs one final refresh after repair is finalized"],
            ["LLM dataset-paper screening", "Not started for full 2020-2025", "Pilot needed", "Cost should be measured before full run"],
            ["Fulltext extraction", "Not started for expanded corpus", "Downstream stage", "Depends on screening positives"],
            ["Metadata enrichment", "Ready after final corpus", "Citation and resource metadata planned", "Run after screening/extraction inputs are fixed"],
        ],
    )

    doc.add_heading("2. Corpus Coverage", level=1)
    add_paragraph(
        doc,
        "The expanded corpus now covers the intended 2020-2025 period across ACL and arXiv. The arXiv screening set is smaller than the repaired corpus because duplicate records and ACL-overlap papers are removed before screening.",
    )
    add_figure(doc, coverage, "Figure 1. Year-level coverage for ACL, repaired arXiv records, and the prepared arXiv screening set.")
    add_table(
        doc,
        ["Year", "ACL catalog", "arXiv repaired corpus", "arXiv screening set"],
        [
            [
                str(year),
                f"{int(acl.get('scope_by_year', {}).get(str(year), 0)):,}",
                f"{int(raw_years.get(str(year), 0)):,}",
                f"{int(screening.get('year_counts', {}).get(str(year), 0)):,}",
            ]
            for year in range(2020, 2026)
        ],
    )

    doc.add_heading("3. arXiv Preparation Summary", level=1)
    add_paragraph(
        doc,
        "The current arXiv screening set removes duplicate titles and papers already represented in the ACL corpus. This prevents the next screening stage from spending model budget on duplicate paper records.",
    )
    add_figure(doc, funnel, "Figure 2. arXiv preparation from repaired records to screening set.")
    add_table(
        doc,
        ["Metric", "Value"],
        [
            ["Input repaired arXiv records", f"{screening.get('raw_rows', 0):,}"],
            ["Unique titles after arXiv deduplication", f"{screening.get('unique_titles_after_arxiv_dedup', 0):,}"],
            ["Removed duplicate titles", f"{screening.get('removed_duplicate_titles', 0):,}"],
            ["ACL paper titles used for overlap removal", f"{screening.get('acl_title_keys', 0):,}"],
            ["Removed ACL title overlap", f"{screening.get('removed_acl_title_overlap', 0):,}"],
            ["Current screening set size", f"{screening.get('output_rows', 0):,}"],
        ],
    )

    doc.add_heading("4. Remaining Issue", level=1)
    add_paragraph(
        doc,
        "The only remaining corpus-preparation issue is a small number of arXiv API timeouts in older 2021 windows. These affect a narrow part of the arXiv repair pass and should be resolved or documented before treating the arXiv corpus as final.",
    )
    add_figure(doc, residual, "Figure 3. Remaining arXiv timeout windows after targeted repair.")
    add_table(
        doc,
        ["Residual window", "Area", "Proposed treatment"],
        [[f"{e['from']} to {e['to']}", e["category"], "Retry at smaller intervals; document if the API continues to fail"] for e in errors],
    )

    doc.add_heading("5. Pipeline Status", level=1)
    add_figure(doc, status, "Figure 4. Current pipeline status before full LLM screening.")

    doc.add_heading("6. Recommended Next Steps", level=1)
    add_table(
        doc,
        ["Step", "Action", "Decision point"],
        [
            ["1", "Finalize the arXiv corpus by incorporating the additional recovered records.", "Check how many new papers are added."],
            ["2", "Refresh the prepared arXiv screening set.", "Confirm duplicate and ACL-overlap removal."],
            ["3", "Retry the remaining timeout windows at smaller intervals.", "If they still fail, document the residual API limitation."],
            ["4", "Run a small LLM screening pilot.", "Estimate cost per 1,000 papers before full run."],
            ["5", "Scale dataset-paper screening only after cost is measured.", "Full screening cost depends on selected model and input size."],
            ["6", "Run fulltext extraction only on screened positives.", "This is the likely highest-cost stage."],
            ["7", "Run metadata enrichment and coverage audit.", "Use public metadata sources first; record query timestamps."],
        ],
    )

    doc.add_heading("7. Cost Estimate", level=1)
    add_table(
        doc,
        ["Stage", "Expected direct cost", "Notes"],
        [
            ["Corpus collection and cleanup", "Near zero", "Local processing and free arXiv/ACL access; cost is runtime."],
            ["Public metadata enrichment", "Near zero to low", "OpenAlex, Semantic Scholar, Hugging Face, GitHub, Papers with Code; rate limits are the main constraint."],
            ["LLM screening pilot", "$5-20", "Recommended before scaling."],
            ["Full LLM screening", "$50-150", "Depends on final paper count, model choice, and input size."],
            ["Fulltext extraction", "$100-300+", "Depends on how many papers pass screening and fulltext length."],
        ],
    )
    add_callout(
        doc,
        "Cost position",
        [
            "A final full-run cost should be set after the pilot.",
            "The next estimate should be cost per 1,000 papers from a representative screening batch.",
            "Fulltext extraction should run only on screened positives to avoid unnecessary spending.",
        ],
    )

    doc.add_heading("8. Summary", level=1)
    add_callout(
        doc,
        "Current interpretation",
        [
            "The corpus expansion is mostly complete at the paper-catalog level.",
            "The main remaining work is to finalize a small remaining arXiv gap, then move to a measured screening pilot.",
            "The cost risk is downstream LLM processing, not corpus collection or public metadata collection.",
        ],
    )

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run(f"Generated {datetime.now().strftime('%Y-%m-%d %H:%M')} | NLP dataset discovery 2020-2025 progress")

    doc.save(OUT_DOCX)
    print(OUT_DOCX)


if __name__ == "__main__":
    main()
